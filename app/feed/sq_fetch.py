# -*- coding: utf-8 -*-
"""SQカレンダー取得 — 自己完結（ベンダリング）。

元: `N225StrategyBuilder/scripts/sq_calendar/generate_sq_calendar.py`（依頼書
   `docs/requests/2026-05-03_SQカレンダー生成.md`）。**import せずコピー**し、保存先を
   LocalEngine 内 `data/sq_calendar.json` / `data/sq_array.pine` に変更（`calendar_fetch` と同型・
   自己完結＝跨ぎ無し）。bs4 のパーサは lxml 非依存の "html.parser" に変更。

JPX 公式の取引最終日 Excel をダウンロードし（現在〜未来年）、過去5年は第2金曜アルゴリズム
（jpholiday）で算出して、メジャーSQ（3/6/9/12月＝OSE:NK225M1! のロール対象）の
「取引最終日の引け時刻」を era 対応で生成する。

出力:
  - `data/sq_calendar.json` … session_lib が読む正本（year_month / last_trading_day /
      session_close_at(ISO・JST) / session_close_unix_ms / regime / sq_date / source）。
  - `data/sq_array.pine`    … KengetsuLib に貼る `session_close_times = array.from(...)` スニペット。

GUI からは `fetch_sq_calendar()` を呼ぶ（別スレッド推奨）。失敗時は既存 JSON を壊さない。

★force_close_* は持たない（引けからのオフセットは利用側＝session_lib / Pine の責務）。
"""
from __future__ import annotations

import calendar  # noqa: F401  (元スクリプト互換・将来利用)
import json
import re
from dataclasses import dataclass
from datetime import datetime, date, time, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Optional

import requests
from bs4 import BeautifulSoup

from .logger import log_message

# ── 依存（self-contained。未導入なら取得は不可だが既存 JSON は使える）──
try:
    import jpholiday
    import openpyxl
    _DEPS_OK = True
    _DEPS_ERR = ""
except Exception as _e:                                    # pragma: no cover
    jpholiday = None  # type: ignore
    openpyxl = None  # type: ignore
    _DEPS_OK = False
    _DEPS_ERR = f"{type(_e).__name__}: {_e}"

JST = timezone(timedelta(hours=9))

# ── 保存先（LocalEngine 内・自己完結）──
DATA_DIR = Path(__file__).resolve().parents[2] / "data"   # = N225LocalEngine/data
JSON_PATH = DATA_DIR / "sq_calendar.json"
PINE_PATH = DATA_DIR / "sq_array.pine"

# ── JPX 取得元（既定。将来 URL 変更時は settings 経由で上書き可）──
DEFAULT_JPX_INDEX_URL = "https://www.jpx.co.jp/derivatives/rules/last-trading-day/index.html"
JPX_EXCEL_URL_TEMPLATE = (
    "https://www.jpx.co.jp/derivatives/rules/last-trading-day/"
    "tvdivq0000004gz8-att/{year}_indexfutures_options_1_j.xlsx"
)
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "ja,en-US;q=0.7,en;q=0.3",
}

MONTH_CODE = {1: "F", 2: "G", 3: "H", 4: "J", 5: "K", 6: "M",
              7: "N", 8: "Q", 9: "U", 10: "V", 11: "X", 12: "Z"}
PAST_YEARS = 5
MAJOR_SQ_MONTHS = (3, 6, 9, 12)               # ラージ・ミニ共通の中心限月切替月

# ── 日中立会終了時刻（クロージング・オークション）の制度履歴（新しい順）──
#   出典: JPX 公式（2024-11-05 東証取引時間延伸に伴い 15:15→15:45）。
#   検索時は「対象日 >= start_date」を満たす最初のエントリを採用。
SESSION_CLOSE_HISTORY: list[tuple[date, time, str]] = [
    (date(2024, 11, 5), time(15, 45), "post-2024-11-05 (TSE extension)"),
    (date(1900, 1, 1), time(15, 15), "pre-2024-11-05 (J-GATE legacy)"),
]


def _session_close_for(d: date) -> tuple[time, str]:
    for start_date, close_time, regime in SESSION_CLOSE_HISTORY:
        if d >= start_date:
            return close_time, regime
    raise RuntimeError(f"No session close regime for {d}")


@dataclass
class _Entry:
    year_month: str
    contract_code: str
    last_trading_day: date
    sq_date: date
    holiday_adjusted: bool
    source: str

    def session_close(self) -> tuple[datetime, str]:
        t, regime = _session_close_for(self.last_trading_day)
        dt = datetime(self.last_trading_day.year, self.last_trading_day.month,
                      self.last_trading_day.day, t.hour, t.minute, tzinfo=JST)
        return dt, regime

    def to_dict(self) -> dict:
        sc, regime = self.session_close()
        return {
            "year_month": self.year_month,
            "contract_code": self.contract_code,
            "last_trading_day": self.last_trading_day.isoformat(),
            "session_close_at": sc.isoformat(),
            "session_close_unix_ms": int(sc.timestamp() * 1000),
            "session_close_regime": regime,
            "sq_date": self.sq_date.isoformat(),
            "holiday_adjusted": self.holiday_adjusted,
            "source": self.source,
        }


# ── 日付ヘルパ ──
def _nth_weekday(year: int, month: int, weekday: int, n: int) -> date:
    first = date(year, month, 1)
    days_to_first = (weekday - first.weekday() + 7) % 7
    return date(year, month, 1 + days_to_first + (n - 1) * 7)


def _next_business_day(d: date) -> date:
    nxt = d + timedelta(days=1)
    while nxt.weekday() >= 5 or jpholiday.is_holiday(nxt):
        nxt += timedelta(days=1)
    return nxt


def _prev_business_day(d: date) -> date:
    prv = d - timedelta(days=1)
    while prv.weekday() >= 5 or jpholiday.is_holiday(prv):
        prv -= timedelta(days=1)
    return prv


def _sq_date(year: int, month: int) -> date:
    """SQ算出日＝第2金曜（祝日なら繰上げ）。"""
    sq = _nth_weekday(year, month, 4, 2)
    while jpholiday.is_holiday(sq):
        sq -= timedelta(days=1)
        while sq.weekday() >= 5:
            sq -= timedelta(days=1)
    return sq


def _make_code(year: int, month: int) -> str:
    return f"{MONTH_CODE[month]}{year % 100:02d}{month:02d}"


def _next_ym(y: int, m: int) -> tuple[int, int]:
    return (y + 1, 1) if m == 12 else (y, m + 1)


def _holiday_adjusted(year: int, month: int, sq: date, last_day: date) -> bool:
    normal_sq = _nth_weekday(year, month, 4, 2)
    exp_last = sq - timedelta(days=1)
    while exp_last.weekday() >= 5 or jpholiday.is_holiday(exp_last):
        exp_last -= timedelta(days=1)
    return (sq != normal_sq) or (last_day != exp_last)


# ── JPX 利用可能年 ──
def _fetch_years(index_url: str, log) -> list[int]:
    log(f"JPX index 取得: {index_url}")
    try:
        resp = requests.get(index_url, headers=HEADERS, timeout=30)
        if resp.status_code != 200:
            raise RuntimeError(f"HTTP {resp.status_code}")
        resp.encoding = "utf-8"
        soup = BeautifulSoup(resp.text, "html.parser")
        pat = re.compile(r"(\d{4})_indexfutures_options_1_j\.xlsx")
        years = {int(m.group(1)) for a in soup.find_all("a", href=True)
                 if (m := pat.search(a["href"]))}
        if not years:
            raise RuntimeError("年リンク抽出失敗")
        out = sorted(years)
        log(f"  年抽出: {out}")
        return out
    except Exception as e:
        log(f"  HTMLパース失敗: {e} → HEAD プローブにフォールバック")
        cur = datetime.now(JST).year
        found = []
        for off in (0, 1, 2):
            y = cur + off
            try:
                r = requests.head(JPX_EXCEL_URL_TEMPLATE.format(year=y), headers=HEADERS,
                                  timeout=15, allow_redirects=True)
                if r.status_code == 200:
                    found.append(y)
            except Exception:
                pass
        if not found:
            raise RuntimeError("JPX Excel を発見できません（サイト仕様変更の可能性）")
        return found


def _download_excel(year: int, log) -> bytes:
    url = JPX_EXCEL_URL_TEMPLATE.format(year=year)
    log(f"  DL: {url}")
    resp = requests.get(url, headers=HEADERS, timeout=60)
    if resp.status_code != 200:
        raise RuntimeError(f"Excel DL 失敗 {year}: HTTP {resp.status_code}")
    return resp.content


# ── Excel パース（日経225先物＝ラージ＝メジャーSQのみ）──
def _is_major_product(s: str) -> bool:
    sl = s.lower()
    if "225" not in s and "２２５" not in s:
        return False
    if "mini" in sl or "ミニ" in s or "マイクロ" in s or "micro" in sl:
        return False
    if "オプション" in s or "option" in sl:
        return False
    return "先物" in s or "futures" in sl


def _identify_columns(ws) -> Optional[dict]:
    product_col = ym_col = last_col = None
    for ri, row in enumerate(ws.iter_rows(values_only=False)):
        if ri > 5:
            break
        for cell in row:
            v = cell.value
            if not isinstance(v, str):
                continue
            vc = v.replace(" ", "").replace("　", "").replace("\n", "")
            if product_col is None and ("商品" in vc or "銘柄" in vc):
                product_col = cell.column
            if ym_col is None and ("取引月" in vc or "限月" in vc):
                ym_col = cell.column
            if last_col is None and "取引最終日" in vc:
                last_col = cell.column
    if product_col and ym_col and last_col:
        return {"product": product_col, "year_month": ym_col, "last_day": last_col}
    return None


_YM_PATS = [re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月"),
            re.compile(r"(\d{4})/(\d{1,2})"), re.compile(r"(\d{4})-(\d{1,2})")]


def _parse_ym_label(value, default_year: int) -> Optional[tuple[int, int]]:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return (value.year, value.month)
    if not isinstance(value, str):
        return None
    s = value.strip()
    for pat in _YM_PATS:
        if (m := pat.search(s)):
            y, mo = int(m.group(1)), int(m.group(2))
            if 1 <= mo <= 12 and 2000 <= y <= 2100:
                return (y, mo)
    if (m := re.search(r"^\s*(\d{1,2})\s*月", s)):
        mo = int(m.group(1))
        if 1 <= mo <= 12:
            return (default_year, mo)
    return None


def _parse_date_cell(value, default_year: int) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        return None
    s = value.strip()
    if (m := re.search(r"(\d{4})[/-](\d{1,2})[/-](\d{1,2})", s)):
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    if (m := re.search(r"^(\d{1,2})/(\d{1,2})$", s)):
        try:
            return date(default_year, int(m.group(1)), int(m.group(2)))
        except ValueError:
            return None
    if (m := re.search(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日", s)):
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    return None


def _extract_from_sheet(ws, year: int, log, sheet_name: str) -> list[_Entry]:
    cols = _identify_columns(ws) or {"product": 2, "year_month": 3, "last_day": 6}
    p, ym, ld = cols["product"], cols["year_month"], cols["last_day"]
    need = max(p, ym, ld)
    out: list[_Entry] = []
    for row in ws.iter_rows(values_only=True):
        if len(row) < need:
            continue
        prod = row[p - 1]
        if not isinstance(prod, str) or not _is_major_product(prod):
            continue
        ymv = _parse_ym_label(row[ym - 1], default_year=year)
        if ymv is None:
            continue
        last_day = _parse_date_cell(row[ld - 1], default_year=ymv[0])
        if last_day is None:
            continue
        if abs((last_day.year - ymv[0]) * 12 + (last_day.month - ymv[1])) > 1:
            continue
        sq = _next_business_day(last_day)
        out.append(_Entry(f"{ymv[0]:04d}-{ymv[1]:02d}", _make_code(ymv[0], ymv[1]),
                          last_day, sq, _holiday_adjusted(ymv[0], ymv[1], sq, last_day),
                          "jpx_excel"))
    log(f"    [{sheet_name}] 日経225先物 抽出: {len(out)} 件")
    return out


def _parse_excel(content: bytes, year: int, log) -> list[_Entry]:
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    entries: list[_Entry] = []
    for sn in wb.sheetnames:
        entries.extend(_extract_from_sheet(wb[sn], year, log, sn))
    seen, uniq = set(), []
    for e in sorted(entries, key=lambda e: e.year_month):
        if e.year_month not in seen:
            seen.add(e.year_month); uniq.append(e)
    return uniq


# ── 過去5年算出 ──
def _calc_past(exec_date: date, jpx_yms: set[str], log) -> list[_Entry]:
    end_y, end_m = (exec_date.year, exec_date.month - 1) if exec_date.month > 1 else (exec_date.year - 1, 12)
    cur_y, cur_m = exec_date.year - PAST_YEARS, exec_date.month
    out: list[_Entry] = []
    while (cur_y, cur_m) <= (end_y, end_m):
        if cur_m in MAJOR_SQ_MONTHS:
            ym = f"{cur_y:04d}-{cur_m:02d}"
            if ym not in jpx_yms:
                sq = _sq_date(cur_y, cur_m)
                last_day = _prev_business_day(sq)
                out.append(_Entry(ym, _make_code(cur_y, cur_m), last_day, sq,
                                  _holiday_adjusted(cur_y, cur_m, sq, last_day), "calculated"))
        cur_y, cur_m = _next_ym(cur_y, cur_m)
    log(f"  過去算出: {len(out)} 件")
    return out


def _integrate(jpx: list[_Entry], past: list[_Entry]) -> list[_Entry]:
    by = {e.year_month: e for e in jpx}
    for e in past:
        by.setdefault(e.year_month, e)
    return sorted(by.values(), key=lambda e: e.year_month)


# ── 出力生成 ──
def _gen_pine(entries: list[_Entry]) -> str:
    lines = ["// ============================================================",
             "// 日経225 メジャー SQ 取引最終日 引け時刻カレンダー（自動生成 by sq_fetch.py）",
             "// 対象: ラージ・ミニ共通の中心限月 (3/6/9/12 月限)",
             "// 値=取引最終日の引け時刻（日中立会終了・制度で変動）。オフセットは Pine 側責務。",
             "// 引け時刻制度（新しい順）:"]
    for sd, t, regime in SESSION_CLOSE_HISTORY:
        lines.append(f"//   {sd.isoformat()} 〜 : {t.strftime('%H:%M')} JST ({regime})")
    lines += ["// ============================================================",
              "session_close_times = array.from("]
    body = []
    for e in entries:
        d = e.last_trading_day; sc, _ = e.session_close()
        body.append(f'    timestamp("Asia/Tokyo", {d.year}, {d.month}, {d.day}, '
                    f'{sc.hour}, {sc.minute}),  // {e.year_month}限 引け {sc.strftime("%H:%M")}'
                    f'{" [祝日繰上げ]" if e.holiday_adjusted else ""}')
    if body:
        body[-1] = body[-1].replace(",  // ", "   // ", 1)
    lines += body + [")", "",
                     "contract_codes = array.from("]
    for i, e in enumerate(entries):
        lines.append(f'    "{e.contract_code}"{"," if i < len(entries) - 1 else ""}')
    lines.append(")")
    return "\n".join(lines) + "\n"


def _gen_json(entries: list[_Entry], exec_date: date, jpx_years: list[int]) -> str:
    now = datetime.now(JST)
    jpx_n = sum(1 for e in entries if e.source == "jpx_excel")
    past_n = sum(1 for e in entries if e.source == "calculated")
    obj = {
        "generated_at": now.isoformat(),
        "execution_date": exec_date.isoformat(),
        "scope": "Major SQ only (3/6/9/12) - Large(NK225)/Mini(NK225M) center contract rolls",
        "coverage": {"first_entry": entries[0].year_month, "last_entry": entries[-1].year_month,
                     "total_entries": len(entries), "past_calculated_count": past_n,
                     "jpx_excel_count": jpx_n},
        "source": {"future_excel_urls": [JPX_EXCEL_URL_TEMPLATE.format(year=y) for y in jpx_years],
                   "past_method": "calculated (2nd Friday of 3/6/9/12, holiday-adjusted jpholiday)"},
        "session_close_history": [{"start_date": d.isoformat(), "close_time_jst": t.strftime("%H:%M"),
                                   "regime_label": label} for d, t, label in SESSION_CLOSE_HISTORY],
        "entries": [e.to_dict() for e in entries],
    }
    return json.dumps(obj, ensure_ascii=False, indent=2)


# ── 公開 API（GUI から呼ぶ）──
def fetch_sq_calendar(index_url: str | None = None) -> dict:
    """JPX から SQ を取得し data/sq_calendar.json + data/sq_array.pine を更新。

    戻り: {"ok":bool, "error":str, "log":[...], "first","last","total","jpx","past",
           "json_path","pine_path"}。失敗時は既存ファイルを壊さない（書かない）。
    """
    logs: list[str] = []

    def log(msg: str):
        line = f"[{datetime.now(JST):%H:%M:%S}] {msg}"
        logs.append(line)
        log_message(f"SQ取得: {msg}", target="info")

    if not _DEPS_OK:
        log(f"依存不足のため取得不可（{_DEPS_ERR}）。jpholiday/openpyxl を導入してください。")
        return {"ok": False, "error": _DEPS_ERR, "log": logs}

    try:
        exec_date = datetime.now(JST).date()
        url = index_url or DEFAULT_JPX_INDEX_URL
        years = _fetch_years(url, log)
        jpx: list[_Entry] = []
        for y in years:
            try:
                jpx.extend(_parse_excel(_download_excel(y, log), y, log))
            except Exception as e:
                log(f"  {y}年 失敗（継続）: {e}")
        if not jpx:
            raise RuntimeError("JPX Excel から1件も取得できませんでした（URL/書式変更を疑う）")
        seen, uniq = set(), []
        for e in sorted(jpx, key=lambda e: e.year_month):
            if e.year_month not in seen:
                seen.add(e.year_month); uniq.append(e)
        jpx = uniq
        past = _calc_past(exec_date, set(seen), log)
        entries = _integrate(jpx, past)

        DATA_DIR.mkdir(parents=True, exist_ok=True)
        PINE_PATH.write_text(_gen_pine(entries), encoding="utf-8")
        JSON_PATH.write_text(_gen_json(entries, exec_date, years), encoding="utf-8")
        log(f"完了: {len(entries)} 件（{entries[0].year_month}〜{entries[-1].year_month}）"
            f" JPX={len(jpx)} 過去={len(past)} → {JSON_PATH.name}/{PINE_PATH.name}")
        return {"ok": True, "error": "", "log": logs,
                "first": entries[0].year_month, "last": entries[-1].year_month,
                "total": len(entries), "jpx": len(jpx), "past": len(past),
                "json_path": str(JSON_PATH), "pine_path": str(PINE_PATH)}
    except Exception as e:
        log(f"取得失敗（既存 JSON は保持）: {type(e).__name__}: {e}")
        return {"ok": False, "error": f"{type(e).__name__}: {e}", "log": logs}


def load_sq_calendar() -> list[dict]:
    """data/sq_calendar.json の entries を返す（無ければ空）。session_lib が読む。"""
    try:
        if JSON_PATH.exists():
            return json.loads(JSON_PATH.read_text(encoding="utf-8")).get("entries", [])
    except Exception as e:
        log_message(f"SQ JSON 読込エラー: {e}", target="info")
    return []


if __name__ == "__main__":
    import sys
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    r = fetch_sq_calendar()
    print(json.dumps({k: v for k, v in r.items() if k != "log"}, ensure_ascii=False, indent=2))
    for ln in r.get("log", []):
        print(ln)
